import io
import json
import uuid
import logging
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader

try:
    import docx
except ImportError:
    docx = None

from app.config import settings
from app.utils.text_splitter import RecursiveCharacterTextSplitter
from app.services.embedding_service import embedding_service
from app.database.pinecone_client import vector_store

logger = logging.getLogger(__name__)

class DocumentService:
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP
        )
        self.docs_dir = Path(__file__).resolve().parent.parent / "data" / "documents"
        self.docs_dir.mkdir(parents=True, exist_ok=True)

    def extract_text_from_bytes(self, content: bytes, filename: str) -> str:
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        
        # PDF Extraction
        if ext == "pdf":
            try:
                reader = PdfReader(io.BytesIO(content))
                pages_text = []
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        pages_text.append(f"--- Page {i + 1} ---\n{text}")
                return "\n\n".join(pages_text)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read PDF file: {str(e)}")
        
        # DOCX Extraction
        elif ext in ["docx", "doc"]:
            if docx is None:
                raise HTTPException(status_code=500, detail="python-docx library is not installed on the server.")
            try:
                doc = docx.Document(io.BytesIO(content))
                elements_text = []
                for p in doc.paragraphs:
                    p_text = p.text.strip()
                    if p_text:
                        elements_text.append(p_text)
                
                # Also extract table text
                for t_idx, table in enumerate(doc.tables):
                    table_rows = []
                    for row in table.rows:
                        row_content = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_content:
                            table_rows.append(row_content)
                    if table_rows:
                        elements_text.append(f"[Table {t_idx + 1}]\n" + "\n".join(table_rows))
                
                return "\n\n".join(elements_text)
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read DOCX file: {str(e)}")

        # Plain Text / Markdown / Code / JSON / CSV
        elif ext in ["txt", "md", "csv", "json", "py", "html", "css", "js", "log"]:
            try:
                return content.decode("utf-8", errors="replace")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read text file: {str(e)}")
        
        else:
            # Fallback: attempt utf-8 decode
            try:
                return content.decode("utf-8", errors="replace")
            except Exception:
                raise HTTPException(status_code=400, detail=f"Unsupported file format: .{ext}")

    async def process_document(self, file: UploadFile) -> Dict[str, Any]:
        filename = file.filename or "uploaded_document"
        content = await file.read()
        file_size = len(content)

        if file_size == 0:
            raise HTTPException(status_code=400, detail="Uploaded file is empty.")

        # Extract text
        raw_text = self.extract_text_from_bytes(content, filename)
        if not raw_text.strip():
            raise HTTPException(status_code=400, detail="No readable text found in document.")

        doc_id = str(uuid.uuid4())[:8]
        uploaded_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
        file_ext = filename.split(".")[-1].lower() if "." in filename else "txt"

        doc_metadata = {
            "doc_id": doc_id,
            "filename": filename,
            "file_type": file_ext,
            "file_size": file_size,
            "uploaded_at": uploaded_at,
        }

        # Chunk document
        chunks = self.text_splitter.create_chunks(raw_text, metadata=doc_metadata)
        if not chunks:
            raise HTTPException(status_code=400, detail="Document could not be chunked.")

        # Extract raw texts for embedding
        chunk_texts = [c["text"] for c in chunks]
        embeddings = embedding_service.embed_documents(chunk_texts)

        # Prepare vector objects
        vectors = []
        for chunk, emb in zip(chunks, embeddings):
            vectors.append({
                "id": chunk["id"],
                "values": emb,
                "metadata": chunk["metadata"]
            })

        # Store in vector database
        vector_store.upsert_vectors(vectors)

        # Save document preview data to disk
        preview_data = {
            "doc_id": doc_id,
            "filename": filename,
            "file_type": file_ext,
            "file_size": file_size,
            "uploaded_at": uploaded_at,
            "total_chunks": len(chunks),
            "char_count": len(raw_text),
            "full_text": raw_text,
            "chunks": [
                {
                    "chunk_id": c["id"],
                    "chunk_index": idx,
                    "text": c["text"],
                    "length": len(c["text"])
                }
                for idx, c in enumerate(chunks)
            ]
        }
        self._save_preview_data(doc_id, preview_data)

        return {
            "doc_id": doc_id,
            "filename": filename,
            "file_type": file_ext,
            "total_chunks": len(chunks),
            "file_size": file_size,
            "uploaded_at": uploaded_at,
            "status": "indexed"
        }

    def _save_preview_data(self, doc_id: str, data: Dict[str, Any]):
        try:
            file_path = self.docs_dir / f"{doc_id}.json"
            with open(file_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.error(f"Error saving document preview data for {doc_id}: {e}")

    def get_document_preview(self, doc_id: str) -> Optional[Dict[str, Any]]:
        file_path = self.docs_dir / f"{doc_id}.json"
        if file_path.exists():
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                logger.error(f"Error reading document preview {doc_id}: {e}")

        # Fallback: Reconstruct preview from vector store metadata
        doc_chunks = []
        doc_filename = "Document"
        file_type = "txt"
        file_size = 0
        uploaded_at = ""

        for item in vector_store.local_vectors:
            meta = item.get("metadata", {})
            if meta.get("doc_id") == doc_id:
                doc_filename = meta.get("filename", doc_filename)
                file_type = meta.get("file_type", file_type)
                file_size = meta.get("file_size", file_size)
                uploaded_at = meta.get("uploaded_at", uploaded_at)
                doc_chunks.append({
                    "chunk_id": item.get("id"),
                    "chunk_index": meta.get("chunk_index", 0),
                    "text": meta.get("text", item.get("text", "")),
                    "length": len(meta.get("text", item.get("text", "")))
                })

        if not doc_chunks:
            return None

        doc_chunks.sort(key=lambda x: x["chunk_index"])
        full_text = "\n\n".join(c["text"] for c in doc_chunks)

        return {
            "doc_id": doc_id,
            "filename": doc_filename,
            "file_type": file_type,
            "file_size": file_size,
            "uploaded_at": uploaded_at,
            "total_chunks": len(doc_chunks),
            "char_count": len(full_text),
            "full_text": full_text,
            "chunks": doc_chunks
        }

    def get_document_text(self, doc_id: str) -> Optional[str]:
        preview = self.get_document_preview(doc_id)
        if preview:
            return preview.get("full_text", "")
        return None

    def list_documents(self) -> List[Dict[str, Any]]:
        return vector_store.list_documents()

    def delete_document(self, doc_id: str) -> bool:
        # Delete from vector store
        deleted = vector_store.delete_document(doc_id)
        # Delete preview file
        file_path = self.docs_dir / f"{doc_id}.json"
        if file_path.exists():
            try:
                file_path.unlink()
            except Exception as e:
                logger.error(f"Error removing document file {file_path}: {e}")
        return deleted

    def clear_all_documents(self) -> bool:
        # Clear vector store
        vector_store.clear_all()
        # Clear preview files
        try:
            for f in self.docs_dir.glob("*.json"):
                try:
                    f.unlink()
                except Exception:
                    pass
        except Exception as e:
            logger.error(f"Error clearing documents directory: {e}")
        return True

document_service = DocumentService()
