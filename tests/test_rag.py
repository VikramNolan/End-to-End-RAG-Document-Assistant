import io
import sys
from pathlib import Path

# Add project root and Backend to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
BACKEND_DIR = PROJECT_ROOT / "Backend"
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(BACKEND_DIR))

import docx
from app.services.document_service import document_service
from app.services.embedding_service import embedding_service
from app.services.rag_service import rag_service
from app.database.pinecone_client import vector_store

def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("RAG Assistant Architecture Overview", level=1)
    doc.add_paragraph(
        "Retrieval-Augmented Generation (RAG) is an advanced AI architecture that merges dense semantic "
        "document retrieval with large language model generation to provide context-aware, verifiable answers."
    )
    doc.add_paragraph(
        "Supported file ingestion formats include PDF, DOCX (Microsoft Word), and plain TXT files. "
        "Documents are parsed into semantic chunks using recursive character splitting with configurable overlaps."
    )
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology"
    hdr_cells[2].text = "Description"
    
    row = table.add_row().cells
    row[0].text = "Inference"
    row[1].text = "Groq LPU"
    row[2].text = "Ultra-fast low-latency token generation"

    row2 = table.add_row().cells
    row2[0].text = "Vector DB"
    row2[1].text = "Pinecone"
    row2[2].text = "Serverless cosine similarity indexing"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_rag_pipeline():
    print("=" * 60)
    print("Testing RAG Assistant Pipeline & DOCX Support")
    print("=" * 60)

    # 1. Test DOCX Text Extraction
    print("\n[Step 1] Testing DOCX Text Extraction...")
    docx_bytes = create_sample_docx()
    extracted_text = document_service.extract_text_from_bytes(docx_bytes, "test_doc.docx")
    assert "Retrieval-Augmented Generation" in extracted_text, "DOCX extraction failed!"
    assert "Groq LPU" in extracted_text, "Table extraction from DOCX failed!"
    print(f"-> DOCX Extracted Successfully ({len(extracted_text)} characters)")

    # 2. Test Text Chunking
    print("\n[Step 2] Testing Recursive Character Text Splitter...")
    chunks = document_service.text_splitter.create_chunks(
        extracted_text,
        metadata={"filename": "test_doc.docx", "doc_id": "test_001"}
    )
    assert len(chunks) > 0, "Chunking produced 0 chunks!"
    print(f"-> Created {len(chunks)} text chunks.")
    print(f"-> First chunk snippet: {chunks[0]['text'][:80]}...")

    # 3. Test Embedding Generation
    print("\n[Step 3] Testing FastEmbed Vector Embeddings...")
    chunk_texts = [c["text"] for c in chunks]
    embeddings = embedding_service.embed_documents(chunk_texts)
    assert len(embeddings) == len(chunk_texts), "Embedding count mismatch!"
    assert len(embeddings[0]) == 384, f"Unexpected embedding dimension: {len(embeddings[0])}"
    print(f"-> Generated {len(embeddings)} dense vector embeddings (Dimension: {len(embeddings[0])})")

    # 4. Test Upsert to Vector Store
    print(f"\n[Step 4] Testing Vector Store Upsert (Mode: {vector_store.client_mode})...")
    vectors = []
    for c, emb in zip(chunks, embeddings):
        vectors.append({
            "id": c["id"],
            "values": emb,
            "metadata": c["metadata"]
        })
    vector_store.upsert_vectors(vectors)
    print("-> Upsert succeeded.")

    # 5. Test Query Retrieval
    print("\n[Step 5] Testing Semantic Query Retrieval...")
    query = "What technologies are used for inference and vector database in RAG Assistant?"
    matches = rag_service.retrieve_context(query, top_k=2)
    assert len(matches) > 0, "No matches retrieved!"
    top_match = matches[0]
    print(f"-> Top retrieved chunk (Score: {top_match['score']}):")
    print(f"   \"{top_match['text'][:120]}...\"")

    # 6. Test LLM Answer Generation
    print("\n[Step 6] Testing Groq LLM Synthesizer...")
    result = rag_service.answer_query(query)
    print(f"-> Model: {result['model']}")
    print(f"-> Total Sources: {result['total_sources_retrieved']}")
    print(f"-> Generated Answer:\n{result['answer']}")

    print("\n" + "=" * 60)
    print("ALL TESTS PASSED SUCCESSFULLY! RAG SYSTEM IS READY.")
    print("=" * 60)

if __name__ == "__main__":
    test_rag_pipeline()
